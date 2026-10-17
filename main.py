"""
E-commerce Price & Stock Tracker with price history and Excel export (USD).

- Crawls a product catalog with pagination (practice site books.toscrape.com, prices in GBP).
- Converts prices to USD with a configurable exchange rate (GBP_USD_RATE).
- Saves every run to a CSV history and calculates the change vs the previous run.
- Exports an Excel report: USD formatting, ratings, product links, conditional formatting,
  and a summary sheet with key stats and a price-distribution chart.

NOTE: the "List Price" is SIMULATED (random 10%-20% markup, stable per product),
because the practice site does not publish real list prices.
"""

import math
import os
import random
import re
import time
from datetime import datetime
from urllib.parse import urljoin

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ----------------------------- Configuration -----------------------------
# Whole catalog (50 pages, ~1,000 products). For a single category use e.g.:
#   http://books.toscrape.com/catalogue/category/books/travel_2/index.html
CATALOG_URL = os.getenv("CATALOG_URL", "http://books.toscrape.com/catalogue/page-1.html")
MAX_PAGES = int(os.getenv("MAX_PAGES", "0")) or None   # 0 / unset = all pages
REQUEST_DELAY = 0.5               # seconds between requests (be polite to the site)
OUTPUT_FILE = "price_monitor_report.xlsx"
HISTORY_FILE = "price_history.csv"

# GBP -> USD exchange rate. Update it before every real run:
#   Linux/Mac: GBP_USD_RATE=1.34 python main.py
GBP_USD_RATE = float(os.getenv("GBP_USD_RATE", "1.34"))

SIM_MARKUP_RANGE = (0.10, 0.20)   # simulated markup for the list price
HIGHLIGHT_DISCOUNT = 0.13         # savings >= 13% are highlighted in green
REQUEST_TIMEOUT = 15              # seconds
USD_FORMAT = '"$"#,##0.00'        # shows "$" regardless of regional settings
PCT_FORMAT = "0.0%"

HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
RATING_WORDS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}


# ------------------------------ Extraction -------------------------------
def fetch_products():
    """Crawls the catalog (with pagination) and returns a list of products."""
    session = requests.Session()
    session.headers.update(HEADERS)
    products, url, pages = [], CATALOG_URL, 0

    while url and (MAX_PAGES is None or pages < MAX_PAGES):
        try:
            response = session.get(url, timeout=REQUEST_TIMEOUT)
            response.raise_for_status()
        except requests.RequestException as exc:
            print(f"❌ Could not connect to the store ({url}): {exc}")
            break

        response.encoding = "utf-8"  # avoids broken characters (â\x80\x99, Â£)
        soup = BeautifulSoup(response.text, "html.parser")

        for item in soup.select("article.product_pod"):
            title = item.h3.a["title"]
            price_text = item.select_one("p.price_color").text
            price_gbp = float(re.search(r"\d+(?:\.\d+)?", price_text).group())
            availability = item.select_one("p.instock.availability").text.strip()
            rating_tag = item.select_one("p.star-rating")
            rating_word = next((c for c in (rating_tag["class"] if rating_tag else []) if c in RATING_WORDS), None)
            products.append(
                {
                    "title": title,
                    "price_gbp": price_gbp,
                    "availability": availability,
                    "rating": RATING_WORDS.get(rating_word),
                    "url": urljoin(url, item.h3.a["href"]),
                }
            )

        pages += 1
        print(f"  page {pages}: {len(products)} products so far")
        next_link = soup.select_one("li.next a")
        url = urljoin(url, next_link["href"]) if next_link else None
        if url:
            time.sleep(REQUEST_DELAY)

    return products


# ------------------------------- Analysis --------------------------------
def simulated_list_price(title, offer_price_usd):
    """Simulated list price: 10%-20% markup, stable per product (does not change between runs)."""
    markup = random.Random(title).uniform(*SIM_MARKUP_RANGE)
    return round(offer_price_usd * (1 + markup), 2)


def build_dataframe(products):
    rows = []
    for p in products:
        offer_usd = round(p["price_gbp"] * GBP_USD_RATE, 2)
        list_usd = simulated_list_price(p["title"], offer_usd)
        rows.append(
            {
                "Product": p["title"],
                "Rating (1-5)": p["rating"],
                "List Price - Simulated ($)": list_usd,
                "Offer Price ($)": offer_usd,
                "Savings ($)": round(list_usd - offer_usd, 2),
                "Savings (%)": round((list_usd - offer_usd) / list_usd, 4),  # fraction: 0.13 = 13%
                "Stock Status": p["availability"],
                "Change vs Previous Run (%)": None,
                "Product URL": p["url"],
                "_price_gbp": p["price_gbp"],
            }
        )
    return pd.DataFrame(rows)


def apply_history(df, run_ts):
    """Compares against the last saved run and appends the current run to the history."""
    if os.path.exists(HISTORY_FILE):
        hist = pd.read_csv(HISTORY_FILE)
        if "product" not in hist.columns:
            print("⚠️ Old-format history file detected; delete it to start a fresh history.")
        else:
            last = hist.sort_values("timestamp").groupby("product").tail(1).set_index("product")
            prev_gbp = df["Product"].map(last["price_gbp"])
            df["Change vs Previous Run (%)"] = ((df["_price_gbp"] - prev_gbp) / prev_gbp).round(4)

    snapshot = pd.DataFrame(
        {
            "timestamp": run_ts,
            "product": df["Product"],
            "price_gbp": df["_price_gbp"],
            "exchange_rate": GBP_USD_RATE,
            "offer_price_usd": df["Offer Price ($)"],
            "stock": df["Stock Status"],
        }
    )
    snapshot.to_csv(HISTORY_FILE, mode="a", header=not os.path.exists(HISTORY_FILE),
                    index=False, encoding="utf-8")
    return df.drop(columns=["_price_gbp"])


# ------------------------------- Export ----------------------------------
def price_distribution(df, step=10):
    """Counts products per price range ($0-9.99, $10-19.99, ...)."""
    top = max(step, math.ceil(df["Offer Price ($)"].max() / step) * step)
    edges = list(range(0, top + step, step))
    labels = [f"${lo}-{hi - 0.01:.2f}" for lo, hi in zip(edges[:-1], edges[1:])]
    counts = pd.cut(df["Offer Price ($)"], bins=edges, labels=labels, right=False).value_counts(sort=False)
    return pd.DataFrame({"Price range": labels, "Products": [int(counts[l]) for l in labels]})


def export_price_analytics(df, run_ts, filename=OUTPUT_FILE):
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center", vertical="center")

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Price Monitor")
        ws = writer.sheets["Price Monitor"]
        col = {name: i + 1 for i, name in enumerate(df.columns)}  # name -> column number
        last_row = len(df) + 1

        for cell in ws[1]:
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        usd_cols = ["List Price - Simulated ($)", "Offer Price ($)", "Savings ($)"]
        pct_cols = ["Savings (%)", "Change vs Previous Run (%)"]
        center_cols = ["Rating (1-5)", "Stock Status"]
        link_font = Font(color="0563C1", underline="single")
        for row in ws.iter_rows(min_row=2):
            for name in usd_cols:
                row[col[name] - 1].number_format = USD_FORMAT
            for name in pct_cols:
                row[col[name] - 1].number_format = PCT_FORMAT
                row[col[name] - 1].alignment = center
            for name in center_cols:
                row[col[name] - 1].alignment = center
            url_cell = row[col["Product URL"] - 1]
            if url_cell.value:
                url_cell.hyperlink = url_cell.value
                url_cell.font = link_font

        # Green: savings >= threshold
        letter = get_column_letter(col["Savings (%)"])
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{last_row}",
            CellIsRule(operator="greaterThanOrEqual", formula=[str(HIGHLIGHT_DISCOUNT)],
                       fill=PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
                       font=Font(color="006100", bold=True)),
        )
        # Red: price went up vs the previous run
        letter = get_column_letter(col["Change vs Previous Run (%)"])
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{last_row}",
            CellIsRule(operator="greaterThan", formula=["0"],
                       fill=PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
                       font=Font(color="9C0006", bold=True)),
        )

        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        for c in ws.columns:
            sample = [len(str(x.value or "")) for x in c[:200]]  # sample rows: fast on big reports
            ws.column_dimensions[get_column_letter(c[0].column)].width = min(max(max(sample) + 4, 14), 60)

        # ---- Summary sheet: key stats + price distribution chart ----
        stats = [
            ("Run date", run_ts, None),
            ("Source", CATALOG_URL, None),
            ("Source currency", "GBP (£)", None),
            ("Report currency", "USD ($)", None),
            ("Exchange rate GBP→USD", GBP_USD_RATE, "0.0000"),
            ("Products monitored", len(df), "#,##0"),
            ("Average offer price ($)", round(df["Offer Price ($)"].mean(), 2), USD_FORMAT),
            ("Lowest offer price ($)", df["Offer Price ($)"].min(), USD_FORMAT),
            ("Highest offer price ($)", df["Offer Price ($)"].max(), USD_FORMAT),
            ("Average savings (%)", round(df["Savings (%)"].mean(), 4), PCT_FORMAT),
            ("Highest savings (%)", df["Savings (%)"].max(), PCT_FORMAT),
            ("Products with price increase", int((df["Change vs Previous Run (%)"] > 0).sum()), "#,##0"),
            ("Products with price drop", int((df["Change vs Previous Run (%)"] < 0).sum()), "#,##0"),
            ("Note", "List price is simulated (10%-20% markup). Prices exclude taxes and shipping.", None),
        ]
        pd.DataFrame({"Item": [s[0] for s in stats], "Value": [s[1] for s in stats]}).to_excel(
            writer, index=False, sheet_name="Summary")
        ws2 = writer.sheets["Summary"]
        for i, (_, _, fmt) in enumerate(stats, start=2):
            cell = ws2.cell(row=i, column=2)
            cell.alignment = Alignment(horizontal="left", wrap_text=True)
            if fmt:
                cell.number_format = fmt
        ws2.column_dimensions["A"].width = 32
        ws2.column_dimensions["B"].width = 60

        dist = price_distribution(df)
        ws2["D1"], ws2["E1"] = "Price range", "Products"
        for r, (label, count) in enumerate(zip(dist["Price range"], dist["Products"]), start=2):
            ws2.cell(row=r, column=4, value=label)
            ws2.cell(row=r, column=5, value=count)
        ws2.column_dimensions["D"].width = 20
        ws2.column_dimensions["E"].width = 12
        for cell in list(ws2[1][:2]) + [ws2["D1"], ws2["E1"]]:
            cell.fill, cell.font = header_fill, header_font

        chart = BarChart()
        chart.type, chart.style = "col", 10
        chart.title = "Products by offer price range ($)"
        chart.y_axis.title, chart.x_axis.title = "Products", "Price range"
        chart.add_data(Reference(ws2, min_col=5, min_row=1, max_row=len(dist) + 1), titles_from_data=True)
        chart.set_categories(Reference(ws2, min_col=4, min_row=2, max_row=len(dist) + 1))
        chart.legend = None
        chart.x_axis.delete = False  # keep axes visible in recent Excel versions
        chart.y_axis.delete = False
        chart.height, chart.width = 8.5, 20
        ws2.add_chart(chart, "G2")

    print(f"✅ Report exported to {filename}")


if __name__ == "__main__":
    print("Starting price monitor...")
    products = fetch_products()
    if products:
        run_ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        frame = apply_history(build_dataframe(products), run_ts)
        export_price_analytics(frame, run_ts)
    else:
        print("No products found; report not generated.")
